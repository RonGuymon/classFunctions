"""nb2docx — convert a Jupyter notebook (.ipynb) to a Word document (.docx).

Python API:
    from class_functions import nb2docx
    nb2docx("notebook.ipynb")

CLI (after pip install):
    nb2docx notebook.ipynb
    nb2docx notebook.ipynb --syntax-highlighting tango
    nb2docx notebook.ipynb --keep-intermediates
    nb2docx notebook.ipynb --regen-reference

Generated support files (reference.docx, jupyterlab.theme) are stored in the
user config directory so they survive package upgrades without needing write
access to site-packages.

External requirements (not pip-installable):
    pandoc  — brew install pandoc  / https://pandoc.org/installing.html
    jupyter — already present if you are running notebooks
"""

import argparse
import copy
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

from platformdirs import user_config_dir

# Writable per-user directory; created on first run.
CONFIG_DIR = Path(user_config_dir("nb2docx"))
REFERENCE_DOCX = CONFIG_DIR / "reference.docx"
THEME_PATH = CONFIG_DIR / "jupyterlab.theme"

# Colors sourced from JupyterLab light theme CSS (--jp-mirror-editor-* variables).
# bold/italic follow CodeMirror 6 defaults for each token class.
_JUPYTERLAB_THEME = {
    "text-color": "#000000",
    "background-color": None,
    "line-number-color": "#aaaaaa",
    "line-number-background-color": None,
    "text-styles": {
        "Keyword":        {"text-color": "#008000", "background-color": None, "bold": False, "italic": False, "underline": False},
        "ControlFlow":    {"text-color": "#008000", "background-color": None, "bold": False, "italic": False, "underline": False},
        "Import":         {"text-color": "#008000", "background-color": None, "bold": False, "italic": False, "underline": False},
        "BuiltIn":        {"text-color": "#008000", "background-color": None, "bold": False, "italic": False, "underline": False},
        "Other":          {"text-color": "#008000", "background-color": None, "bold": False, "italic": False, "underline": False},
        "DecVal":         {"text-color": "#008800", "background-color": None, "bold": False, "italic": False, "underline": False},
        "BaseN":          {"text-color": "#008800", "background-color": None, "bold": False, "italic": False, "underline": False},
        "Float":          {"text-color": "#008800", "background-color": None, "bold": False, "italic": False, "underline": False},
        "Operator":       {"text-color": "#7800c2", "background-color": None, "bold": False, "italic": False, "underline": False},
        "Comment":        {"text-color": "#408080", "background-color": None, "bold": False, "italic": True,  "underline": False},
        "CommentVar":     {"text-color": "#408080", "background-color": None, "bold": False, "italic": True,  "underline": False},
        "Annotation":     {"text-color": "#408080", "background-color": None, "bold": False, "italic": True,  "underline": False},
        "Information":    {"text-color": "#408080", "background-color": None, "bold": False, "italic": True,  "underline": False},
        "Warning":        {"text-color": "#408080", "background-color": None, "bold": True,  "italic": True,  "underline": False},
        "String":         {"text-color": "#ba2121", "background-color": None, "bold": False, "italic": False, "underline": False},
        "VerbatimString": {"text-color": "#ba2121", "background-color": None, "bold": False, "italic": False, "underline": False},
        "Char":           {"text-color": "#ba2121", "background-color": None, "bold": False, "italic": False, "underline": False},
        "SpecialChar":    {"text-color": "#ba2121", "background-color": None, "bold": False, "italic": False, "underline": False},
        "Documentation":  {"text-color": "#ba2121", "background-color": None, "bold": False, "italic": True,  "underline": False},
        "SpecialString":  {"text-color": "#770088", "background-color": None, "bold": False, "italic": False, "underline": False},
        "Function":       {"text-color": "#0000ff", "background-color": None, "bold": False, "italic": False, "underline": False},
        "Variable":       {"text-color": "#212121", "background-color": None, "bold": False, "italic": False, "underline": False},
        "Constant":       {"text-color": "#8888ff", "background-color": None, "bold": False, "italic": False, "underline": False},
        "Preprocessor":   {"text-color": "#aa22ff", "background-color": None, "bold": False, "italic": False, "underline": False},
        "Attribute":      {"text-color": "#0000cc", "background-color": None, "bold": False, "italic": False, "underline": False},
        "DataType":       {"text-color": "#555555", "background-color": None, "bold": False, "italic": False, "underline": False},
        "Error":          {"text-color": "#ff0000", "background-color": None, "bold": True,  "italic": False, "underline": False},
        "Alert":          {"text-color": "#ff0000", "background-color": None, "bold": True,  "italic": False, "underline": False},
        "Extension":      {"text-color": None,      "background-color": None, "bold": False, "italic": False, "underline": False},
    },
}


def _create_theme_file() -> None:
    CONFIG_DIR.mkdir(parents=True, exist_ok=True)
    THEME_PATH.write_text(json.dumps(_JUPYTERLAB_THEME, indent=4))
    print(f"Created: {THEME_PATH}")


def _create_reference_docx() -> None:
    """Build reference.docx from pandoc defaults and add a gray Source Code style."""
    CONFIG_DIR.mkdir(parents=True, exist_ok=True)

    result = subprocess.run(
        ["pandoc", "--print-default-data-file", "reference.docx"],
        capture_output=True,
        check=True,
    )
    REFERENCE_DOCX.write_bytes(result.stdout)

    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document(REFERENCE_DOCX)

    # Pandoc maps fenced code blocks to the "Source Code" paragraph style.
    # The pandoc default reference.docx omits this style, so we add it here
    # with gray shading so pandoc adopts it instead of creating a plain default.
    style = next((s for s in doc.styles if s.name == "Source Code"), None)
    if style is None:
        style = doc.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)

    style.font.name = "Courier New"
    style.font.size = Pt(10)

    pPr = style.element.get_or_add_pPr()

    for el in pPr.findall(qn("w:shd")):
        pPr.remove(el)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "ECECEC")
    pPr.append(shd)

    for el in pPr.findall(qn("w:spacing")):
        pPr.remove(el)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "60")
    spacing.set(qn("w:after"), "60")
    pPr.append(spacing)

    doc.save(REFERENCE_DOCX)
    print(f"Created: {REFERENCE_DOCX}")
    _create_theme_file()


_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_METHOD_PATTERN = re.compile(r"(?<=\.)[a-zA-Z_][a-zA-Z0-9_]*")
_METHOD_COLOR = "0000FF"


def _apply_method_colors(docx_path: Path) -> None:
    """Color identifiers that follow a '.' blue in every Source Code paragraph.

    Pandoc's Python tokenizer lumps method/attribute names into NormalTok, so
    the theme file cannot reach them. This function splits those runs at each
    '.name' boundary and sets the color directly on the name fragment.
    """
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document(docx_path)
    modified = False

    for para in doc.paragraphs:
        if para.style.name != "Source Code":
            continue

        for run_el in list(para._element.findall(qn("w:r"))):
            rpr = run_el.find(qn("w:rPr"))
            if rpr is None:
                continue
            style_el = rpr.find(qn("w:rStyle"))
            if style_el is None or style_el.get(qn("w:val")) != "NormalTok":
                continue

            t_el = run_el.find(qn("w:t"))
            if t_el is None:
                continue
            text = t_el.text or ""

            matches = list(_METHOD_PATTERN.finditer(text))
            if not matches:
                continue

            # Decompose text into (segment, is_method_name) pairs
            segments: list[tuple[str, bool]] = []
            last = 0
            for m in matches:
                if m.start() > last:
                    segments.append((text[last:m.start()], False))
                segments.append((m.group(), True))
                last = m.end()
            if last < len(text):
                segments.append((text[last:], False))

            parent = run_el.getparent()
            idx = list(parent).index(run_el)
            parent.remove(run_el)

            insert_at = idx
            for seg, is_method in segments:
                if not seg:
                    continue
                new_run = copy.deepcopy(run_el)
                new_t = new_run.find(qn("w:t"))
                new_t.text = seg
                if seg != seg.strip():
                    new_t.set(_XML_SPACE, "preserve")
                elif _XML_SPACE in new_t.attrib:
                    del new_t.attrib[_XML_SPACE]

                if is_method:
                    new_rpr = new_run.find(qn("w:rPr"))
                    # Keep NormalTok so the run inherits Consolas 11pt from
                    # VerbatimChar. A direct w:color on the run overrides the
                    # NormalTok color (#000000) without touching the font.
                    color_el = OxmlElement("w:color")
                    color_el.set(qn("w:val"), _METHOD_COLOR)
                    new_rpr.insert(0, color_el)

                parent.insert(insert_at, new_run)
                insert_at += 1
                modified = True

    if modified:
        doc.save(docx_path)


def _set_cell_shading(cell, fill_hex: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = cell._tc.get_or_add_tcPr()
    for existing in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def _style_tables(docx_path: Path) -> None:
    """Style all Word tables produced from DataFrame HTML:
      - Header row (row 0): bold
      - Index column (col 0): bold
      - Every cell: right-aligned
      - Rows 2, 4, 6 … (second data row onward, alternating): light gray fill
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    GRAY = "ECECEC"

    doc = Document(docx_path)
    if not doc.tables:
        return

    for table in doc.tables:
        for row_idx, row in enumerate(table.rows):
            is_header = row_idx == 0
            shade = row_idx >= 2 and row_idx % 2 == 0

            for col_idx, cell in enumerate(row.cells):
                bold = is_header or col_idx == 0

                if shade:
                    _set_cell_shading(cell, GRAY)

                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    for run in para.runs:
                        if bold:
                            run.font.bold = True

    doc.save(docx_path)


def _html_table_to_pipe(html_fragment: str) -> str:
    """Convert HTML table fragment(s) to pipe-table markdown.

    Handles pandas DataFrames (thead/tbody structure) and statsmodels simpletables
    (bare tr rows with mixed th/td, optional caption, no thead/tbody).
    Multiple tables in one fragment are each converted and joined with blank lines.
    """
    from lxml import html as lhtml

    root = lhtml.fragment_fromstring(html_fragment, create_parent="div")
    tables = root.findall(".//table")
    if not tables:
        return html_fragment

    def cell_text(el) -> str:
        return (el.text_content() or "").strip().replace("|", "\\|")

    def table_to_pipe(table) -> str:
        caption_el = table.find("caption")
        caption = caption_el.text_content().strip() if caption_el is not None else ""

        thead = table.find(".//thead")
        tbody = table.find(".//tbody")

        if thead is not None:
            header_rows = thead.findall(".//tr")
            data_rows = tbody.findall("tr") if tbody is not None else []
            headers = [cell_text(c) for c in header_rows[0]] if header_rows else []
        else:
            all_rows = table.findall(".//tr")
            if not all_rows:
                return ""
            first_row = all_rows[0]
            th_count = len(first_row.findall("th"))
            td_count = len(first_row.findall("td"))
            if th_count > td_count:
                headers = [cell_text(c) for c in first_row]
                data_rows = all_rows[1:]
            else:
                headers = []
                data_rows = list(all_rows)

        n = len(headers) if headers else (len(data_rows[0]) if data_rows else 0)
        if n == 0:
            return ""

        if not headers:
            headers = [""] * n

        rows = [[cell_text(c) for c in tr] for tr in data_rows]
        rows = [r[:n] + [""] * max(0, n - len(r)) for r in rows]

        col_widths = [max(3, len(h)) for h in headers]
        for row in rows:
            for i, cell in enumerate(row[:n]):
                col_widths[i] = max(col_widths[i], len(cell))

        def fmt(cells: list[str]) -> str:
            return "| " + " | ".join(c.ljust(w) for c, w in zip(cells, col_widths)) + " |"

        sep = "| " + " | ".join("-" * w for w in col_widths) + " |"
        lines = []
        if caption:
            lines.append(f"**{caption}**\n")
        lines += [fmt(headers), sep] + [fmt(row) for row in rows]
        return "\n".join(lines) + "\n"

    parts = [table_to_pipe(t) for t in tables]
    result = "\n\n".join(p for p in parts if p)
    return result if result else html_fragment


def _html_tables_to_markdown(md: str) -> str:
    """Replace HTML table blocks with pipe-format markdown tables.

    Pass 1: pandas DataFrames wrapped in <div><style>…</style><table>…</table></div>
    Pass 2: bare statsmodels simpletables (no div wrapper)
    Also converts trailing <br/> tags to newlines.
    """
    def replace(match: re.Match) -> str:
        try:
            return _html_table_to_pipe(match.group(0))
        except Exception:
            return match.group(0)

    # Pass 1: div-wrapped tables (pandas)
    div_pattern = re.compile(
        r"<div[^>]*>\s*(?:<style[^>]*>.*?</style>\s*)?<table[^>]*>.*?</table>\s*</div>",
        re.DOTALL | re.IGNORECASE,
    )
    md = div_pattern.sub(replace, md)

    # Pass 2: bare tables (statsmodels simpletables)
    bare_pattern = re.compile(r"<table[^>]*>.*?</table>", re.DOTALL | re.IGNORECASE)
    md = bare_pattern.sub(replace, md)

    # Convert <br/> to newlines (statsmodels trailing notes)
    md = re.sub(r"<br\s*/?>", "\n", md, flags=re.IGNORECASE)

    return md


def nb2docx(
    notebook: "str | Path",
    syntax_highlighting: "str | None" = None,
    keep_intermediates: bool = False,
    regen_reference: bool = False,
) -> Path:
    """Convert a Jupyter notebook to a Word document.

    Parameters
    ----------
    notebook:
        Path to the .ipynb file.
    syntax_highlighting:
        Pandoc syntax highlight style name or path to a .theme file.
        Defaults to the bundled JupyterLab light theme.
    keep_intermediates:
        If True, keep the intermediate .md and _files/ directory.
    regen_reference:
        If True, regenerate reference.docx and jupyterlab.theme even if
        they already exist in the user config directory.

    Returns
    -------
    Path to the generated .docx file.
    """
    notebook_path = Path(notebook).resolve()

    if regen_reference:
        REFERENCE_DOCX.unlink(missing_ok=True)
        THEME_PATH.unlink(missing_ok=True)

    if not REFERENCE_DOCX.exists() or not THEME_PATH.exists():
        print("Creating reference.docx and jupyterlab.theme ...")
        _create_reference_docx()

    highlighting = syntax_highlighting or str(THEME_PATH)
    out_dir = notebook_path.parent
    stem = notebook_path.stem
    md_path = out_dir / f"{stem}.md"
    files_dir = out_dir / f"{stem}_files"
    docx_path = out_dir / f"{stem}.docx"

    print(f"Step 1: {notebook_path.name} → {md_path.name}")
    subprocess.run(
        [
            "jupyter", "nbconvert",
            "--to", "markdown",
            "--output", stem,
            "--output-dir", str(out_dir),
            str(notebook_path),
        ],
        check=True,
    )

    # Pre-process markdown before pandoc:
    #   1. Convert HTML DataFrame tables to pipe tables (pandoc drops raw HTML in docx mode)
    #   2. Strip "png" alt text that nbconvert adds to every image (pandoc renders it as a caption)
    md = md_path.read_text()
    md = _html_tables_to_markdown(md)
    md = re.sub(r"!\[png\]", "![]", md)
    md_path.write_text(md)

    print(f"Step 2: {md_path.name} → {docx_path.name}")
    subprocess.run(
        [
            "pandoc", str(md_path),
            "-o", str(docx_path),
            f"--syntax-highlighting={highlighting}",
            f"--reference-doc={REFERENCE_DOCX}",
        ],
        check=True,
        cwd=out_dir,
    )

    print("Step 3: coloring method names")
    _apply_method_colors(docx_path)

    print("Step 4: styling tables")
    _style_tables(docx_path)

    if not keep_intermediates:
        md_path.unlink(missing_ok=True)
        if files_dir.exists():
            shutil.rmtree(files_dir)

    return docx_path


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Convert a Jupyter notebook to a Word document.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Default syntax highlighting matches JupyterLab's light theme.\n"
            "Built-in style names: pygments, tango, kate, espresso, zenburn,\n"
            "                      breezedark, haddock, monochrome\n\n"
            "External dependencies (not installed by pip):\n"
            "  pandoc  — brew install pandoc\n"
            "  jupyter — pip install jupyter"
        ),
    )
    parser.add_argument("notebook", help="Path to .ipynb file")
    parser.add_argument(
        "--syntax-highlighting",
        default=None,
        metavar="STYLE",
        help="Pandoc style name or path to a .theme file (default: jupyterlab.theme)",
    )
    parser.add_argument(
        "--keep-intermediates",
        action="store_true",
        help="Keep the intermediate .md and _files/ directory after conversion",
    )
    parser.add_argument(
        "--regen-reference",
        action="store_true",
        help="Regenerate reference.docx and jupyterlab.theme even if they already exist",
    )
    args = parser.parse_args()

    notebook_path = Path(args.notebook)
    if not notebook_path.exists():
        print(f"Error: {notebook_path} not found", file=sys.stderr)
        sys.exit(1)
    if notebook_path.suffix != ".ipynb":
        print(f"Error: expected a .ipynb file, got '{notebook_path.suffix}'", file=sys.stderr)
        sys.exit(1)

    docx_path = nb2docx(
        notebook_path,
        syntax_highlighting=args.syntax_highlighting,
        keep_intermediates=args.keep_intermediates,
        regen_reference=args.regen_reference,
    )
    print(f"Done: {docx_path}")


if __name__ == "__main__":
    main()
